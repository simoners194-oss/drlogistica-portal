# DR Portal — modulo Mezzi: generatore del seed iniziale (una tantum).
# ---------------------------------------------------------------------------
# Legge "Dettaglio Mezzi DR.xlsx" (fogli Noleggio, Proprieta, Documenti,
# Telepass+ZTL) e i tre JSON estratti dai PDF (libretti, assicurazioni,
# contratti) e produce un JSON nello schema MezziDb (src/lib/mezzi-types.ts)
# da incollare in Mezzi → Parametri → "Importa database (JSON)".
#
# Uso:
#   python genera_seed.py <cartella-estrazioni> <output.json>
# ---------------------------------------------------------------------------
import json
import re
import sys
import unicodedata
from datetime import date, timedelta
from pathlib import Path

import openpyxl

XLSX = (
    r"C:\Users\simon\OneDrive - DR LOGISTICA\Documenti Condivisi - Documenti"
    r"\Operativo\Mezzi DR\Dettaglio Mezzi DR.xlsx"
)
OGGI = "2026-08-26"

AUTO_DIRIGENZIALI = re.compile(
    r"BMW|PORSCHE|AUDI|GLC|GLB|GLE|CAPTUR|T-ROC|TIPO|M550|SPORTAGE|CAYENNE", re.I
)


def norm_targa(t):
    return re.sub(r"[^A-Z0-9]", "", (t or "").upper())


def iso(v):
    if v is None:
        return None
    s = str(v)[:10]
    return s if re.match(r"^\d{4}-\d{2}-\d{2}$", s) else None


def parse_data_it(testo):
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", testo or "")
    if not m:
        return None
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"


def num(v):
    if v is None:
        return None
    try:
        s = str(v).replace(",", ".").strip()
        if not s or s.lower() in ("illimitato", "illimitati", "non indicato"):
            return None
        return float(s)
    except ValueError:
        return None


def pulisci(s):
    if s is None:
        return None
    s = unicodedata.normalize("NFC", str(s)).strip()
    return s or None


def main():
    estr_dir = Path(sys.argv[1])
    out_path = Path(sys.argv[2])

    libretti = {
        norm_targa(r["targa"]): r
        for r in json.loads((estr_dir / "estrazione-libretti.json").read_text("utf-8"))
        if r.get("targa")
    }
    assicurazioni = json.loads(
        (estr_dir / "estrazione-assicurazioni.json").read_text("utf-8")
    )
    contratti_pdf = json.loads(
        (estr_dir / "estrazione-contratti.json").read_text("utf-8")
    )

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    warn = []

    mezzi = {}       # targa -> mezzo
    affidamenti = []
    contratti = []
    scadenze = []
    ztl = []
    seq = {"n": 0}

    def nid(pref):
        seq["n"] += 1
        return f"{pref}-seed-{seq['n']:03d}"

    def stato_da_note(note, appalto):
        n = (note or "").lower()
        if "reso in data" in n or "da rendere" in n:
            return "reso"
        if "sostitutivo" in n:
            return "sostitutivo"
        if "manutenzione" in n or "guasto" in n or "sinistro" in n:
            return "manutenzione"
        if (appalto or "").strip().lower() == "ufficio":
            return "dirigenziale"
        return "operativo"

    def gruppo_da_stato(stato):
        return "operativo" if stato in ("operativo", "manutenzione", "sostitutivo") else "altro"

    # ---------------- Foglio Noleggio ----------------
    ws = wb["Noleggio"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        targa_raw, tipo, modello, appalto, canone, noleggiatore, scad, autista, note, km_mese, extra = (
            (list(row) + [None] * 11)[:11]
        )
        targa = norm_targa(targa_raw or "")
        if not targa:
            if any(v for v in row):
                warn.append(f"Riga Noleggio senza targa saltata: {pulisci(modello)} {canone}")
            continue
        stato = stato_da_note(note, appalto)
        m = {
            "id": targa,
            "targa": targa,
            "gruppo": gruppo_da_stato(stato),
            "stato": stato,
            "proprieta": "noleggio",
            "societa": "DR Logistica srl",
            "tipoMezzo": pulisci(tipo),
            "modello": pulisci(modello),
            "appalto": pulisci(appalto),
            "note": pulisci(note),
        }
        srt = re.search(r"sostitutivo del (\w{7})", note or "", re.I)
        if srt:
            m["sostitutivoDi"] = norm_targa(srt.group(1))
        mezzi[targa] = m

        # Contratto dal foglio (arricchito dopo con i PDF)
        c = {
            "id": nid("ctr"),
            "mezzoId": targa,
            "noleggiatore": pulisci(noleggiatore),
            "canoneMensileEur": num(canone),
            "dataFine": iso(scad),
            "kmMeseInclusi": num(km_mese),
            "extraKmEur": num(extra),
            "attivo": stato != "reso",
        }
        contratti.append(c)

        # Affidamento corrente se l'autista è una persona vera
        a = pulisci(autista)
        if a and "chiesto info" not in a.lower():
            affidamenti.append(
                {
                    "id": nid("aff"),
                    "mezzoId": targa,
                    "autistaNome": a,
                    "appalto": pulisci(appalto),
                    "dal": OGGI,
                    "note": "Import iniziale: data di inizio reale non nota.",
                }
            )

    # ---------------- Foglio Proprietà ----------------
    ws = wb["Propriet\u00e0"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        targa_raw, modello, sede, fornitore, note = (list(row) + [None] * 5)[:5]
        loc = (list(row) + [None] * 6)[5]
        targa = norm_targa(targa_raw or "")
        if not targa:
            continue
        blob = " ".join(str(x) for x in (modello, sede, fornitore, note, loc) if x)
        bl = blob.lower()
        if "rottamato" in bl:
            stato = "rottamato"
        elif "venduto" in bl:
            stato = "venduto"
        elif "rotto" in bl or "fermo" in bl:
            stato = "fermo"
        elif AUTO_DIRIGENZIALI.search(str(modello) or ""):
            stato = "dirigenziale"
        else:
            stato = "operativo"
        m = {
            "id": targa,
            "targa": targa,
            "gruppo": gruppo_da_stato(stato),
            "stato": stato,
            "proprieta": "proprieta",
            "societa": "DR Soluzione Logistica",
            "modello": pulisci(modello),
            "sede": pulisci(sede),
            "note": " · ".join(x for x in (pulisci(note), pulisci(loc)) if x) or None,
        }
        if targa in mezzi:
            prev = mezzi[targa]
            prev["note"] = " · ".join(x for x in (prev.get("note"), m.get("note")) if x) or None
            warn.append(f"{targa}: duplicato nel foglio Proprieta, righe unite")
        else:
            mezzi[targa] = m

    # ---------------- Foglio Documenti (proprietà: assicurazioni + revisioni) --------
    ws = wb["Documenti mezzi di propriet\u00e0"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        targa_raw, tipo_v, assic, _imp, _mens, revisione, _fatt, _libr, note = (
            (list(row) + [None] * 9)[:9]
        )
        targa = norm_targa(targa_raw or "")
        if not targa:
            continue
        if targa not in mezzi:
            mezzi[targa] = {
                "id": targa,
                "targa": targa,
                "gruppo": "altro",
                "stato": "fermo",
                "proprieta": "proprieta",
                "societa": "DR Soluzione Logistica",
                "tipoMezzo": pulisci(tipo_v),
                "note": "Presente solo nel foglio Documenti.",
            }
            warn.append(f"{targa}: presente solo nel foglio Documenti")
        mezzo = mezzi[targa]
        if not mezzo.get("tipoMezzo"):
            mezzo["tipoMezzo"] = pulisci(tipo_v)
        # Assicurazione: data nel testo, altrimenti nota
        s_ass = parse_data_it(str(assic or ""))
        if s_ass:
            scadenze.append(
                {
                    "id": nid("scad"),
                    "mezzoId": targa,
                    "tipo": "assicurazione",
                    "scadenza": s_ass,
                    "descrizione": pulisci(str(assic).split("-")[0]),
                }
            )
        elif assic and "non risulta" in str(assic).lower():
            mezzo["note"] = " · ".join(
                x for x in (mezzo.get("note"), "NON RISULTA ASSICURATO (da verificare)") if x
            )
        # Revisione: la colonna riporta l'ULTIMA revisione → prossima = +2 anni
        r = iso(revisione)
        if r:
            y, mth, d = map(int, r.split("-"))
            try:
                nxt = date(y + 2, mth, d)
            except ValueError:
                nxt = date(y + 2, mth, 28)
            scadenze.append(
                {
                    "id": nid("scad"),
                    "mezzoId": targa,
                    "tipo": "revisione",
                    "scadenza": nxt.isoformat(),
                    "descrizione": f"Calcolata: ultima revisione {r} + 2 anni — VERIFICARE",
                }
            )

    # ---------------- Mezzi presenti SOLO nei libretti ----------------
    # (es. targhe uscite dal foglio Noleggio ma ancora in flotta)
    for targa, lib in libretti.items():
        if targa in mezzi:
            continue
        mezzi[targa] = {
            "id": targa,
            "targa": targa,
            "gruppo": "operativo",
            "stato": "operativo",
            "proprieta": "noleggio",
            "societa": "DR Logistica srl",
            "marca": pulisci(lib.get("marca")),
            "modello": pulisci(lib.get("modello")),
            "note": "Assente dai fogli Excel: creato dal libretto — verificare stato/appalto.",
        }
        warn.append(f"{targa}: assente dagli elenchi Excel, creato dal libretto")

    # ---------------- Foglio Telepass + ZTL ----------------
    ws = wb["Telepass + ZTL (Univex Milano)"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        targa_raw, giro, autista, telepass, appalto, comune, inizio, fine, note = (
            (list(row) + [None] * 9)[:9]
        )
        targa = norm_targa(targa_raw or "")
        tp = pulisci(telepass)
        if targa and targa in mezzi and tp and tp.lower() != "no telepass":
            mezzi[targa]["telepass"] = tp
        if comune:
            p = {
                "id": nid("ztl"),
                "comune": pulisci(comune),
                "stato": "attivo",
                "dal": iso(inizio),
                "al": iso(fine),
                "note": pulisci(note),
            }
            if targa and targa in mezzi:
                p["mezzoId"] = targa
            if p["al"] and p["al"] < OGGI:
                p["stato"] = "scaduto"
            if not p["dal"] and not p["al"]:
                p["note"] = " · ".join(x for x in (p["note"], "date permesso non note") if x)
            ztl.append(p)
            if p.get("al") and p["stato"] == "attivo" and p.get("mezzoId"):
                scadenze.append(
                    {
                        "id": nid("scad"),
                        "mezzoId": p["mezzoId"],
                        "tipo": "ztl",
                        "scadenza": p["al"],
                        "descrizione": f"Permesso ZTL {p['comune']}",
                    }
                )
        # Il giro come nota sul mezzo
        if targa and targa in mezzi and pulisci(giro):
            mezzi[targa]["note"] = " · ".join(
                x for x in (mezzi[targa].get("note"), f"Giri: {pulisci(giro)}") if x
            )
        # Autista dal foglio Telepass se il mezzo non ha già un affidamento
        a = pulisci(autista)
        if (
            targa
            and targa in mezzi
            and a
            and not any(x["mezzoId"] == targa for x in affidamenti)
        ):
            affidamenti.append(
                {
                    "id": nid("aff"),
                    "mezzoId": targa,
                    "autistaNome": a,
                    "appalto": pulisci(appalto) or mezzi[targa].get("appalto"),
                    "dal": OGGI,
                    "note": "Import iniziale (foglio Telepass): data di inizio reale non nota.",
                }
            )
            mezzi[targa].setdefault("appalto", pulisci(appalto))

    # ---------------- Arricchimento dai libretti ----------------
    for targa, lib in libretti.items():
        if targa not in mezzi:
            warn.append(f"{targa}: libretto estratto ma non presente negli elenchi Excel")
            continue
        mz = mezzi[targa]
        mz["marca"] = mz.get("marca") or pulisci(lib.get("marca"))
        if not mz.get("modello") or (lib.get("modello") and len(str(lib["modello"])) > 3):
            mz["modello"] = pulisci(lib.get("modello")) or mz.get("modello")
        if lib.get("frigo"):
            mz["frigo"] = True
        if lib.get("atp"):
            mz["atp"] = True
        mz["libretto"] = {
            k: v
            for k, v in {
                "telaio": pulisci(lib.get("telaio")),
                "immatricolazione": iso(lib.get("dataImmatricolazione")),
                "intestatario": pulisci(lib.get("intestatario")),
                "massaMaxKg": lib.get("massaMax"),
                "massaVuotoKg": lib.get("massaVuotoKg"),
                "portataKg": lib.get("portataKg"),
                "cilindrata": lib.get("cilindrata"),
                "potenzaKw": lib.get("potenzaKw"),
                "alimentazione": pulisci(lib.get("alimentazione")),
                "classeEuro": pulisci(lib.get("classeEuro")),
                "lunghezzaM": lib.get("lunghezzaM"),
                "larghezzaM": lib.get("larghezzaM"),
                "carrozzeria": pulisci(lib.get("carrozzeria")),
                "allestimento": pulisci(lib.get("allestimento")),
                "posti": lib.get("posti"),
            }.items()
            if v is not None
        }
        if lib.get("file"):
            mz["docLibretto"] = f"LIBRETTI/{lib['file']}"

    # ---------------- Arricchimento contratti dai PDF ----------------
    by_targa = {}
    for c in contratti:
        by_targa.setdefault(c["mezzoId"], c)
    for cp in contratti_pdf:
        targa = norm_targa(cp.get("targa") or "")
        if not targa or targa not in mezzi:
            continue
        cliente = (cp.get("cliente") or "").lower()
        if cliente and "dr logistica" not in cliente:
            warn.append(
                f"{targa}: contratto PDF con cliente '{cp.get('cliente')}' (non DR) — non importato"
            )
            continue
        c = by_targa.get(targa)
        if c is None:
            c = {"id": nid("ctr"), "mezzoId": targa, "attivo": True}
            contratti.append(c)
            by_targa[targa] = c
        c.setdefault("noleggiatore", pulisci(cp.get("noleggiatore")))
        if c.get("canoneMensileEur") is None:
            c["canoneMensileEur"] = cp.get("canoneMensileEur")
        if cp.get("canoneIva") and "+iva" in str(cp["canoneIva"]).lower():
            c["canonePiuIva"] = True
        c["dataInizio"] = iso(cp.get("dataInizio")) or c.get("dataInizio")
        if not c.get("dataFine"):
            c["dataFine"] = iso(cp.get("dataFine"))
        if c.get("kmMeseInclusi") is None:
            c["kmMeseInclusi"] = num(cp.get("kmInclusiMese"))
        if c.get("extraKmEur") is None:
            c["extraKmEur"] = cp.get("costoKmExtraEur")
        if cp.get("preavvisoDisdettaGiorni"):
            c["preavvisoDisdettaGiorni"] = cp["preavvisoDisdettaGiorni"]
        fr = cp.get("franchigie") or {}
        if isinstance(fr, dict):
            if fr.get("rca") is not None:
                c["franchigiaRca"] = str(fr["rca"])
            if fr.get("kasko") is not None:
                c["franchigiaKasko"] = str(fr["kasko"])
            if fr.get("furtoIncendio") is not None:
                c["franchigiaFurto"] = str(fr["furtoIncendio"])
        if cp.get("serviziInclusi"):
            c["serviziInclusi"] = cp["serviziInclusi"]
        if cp.get("penali"):
            c["penali"] = str(cp["penali"])[:500]
        if cp.get("deposito"):
            c["deposito"] = cp["deposito"]
        if cp.get("note"):
            c["note"] = str(cp["note"])[:500]
        if cp.get("file"):
            mezzi[targa].setdefault("docContratto", f"CONTRATTI DI NOLEGGIO/{cp['file']}")

    # Scadenza fine contratto (+ finestra disdetta se c'è preavviso)
    for c in contratti:
        if c.get("attivo") and c.get("dataFine"):
            scadenze.append(
                {
                    "id": nid("scad"),
                    "mezzoId": c["mezzoId"],
                    "tipo": "fine_contratto",
                    "scadenza": c["dataFine"],
                    "descrizione": c.get("noleggiatore"),
                }
            )
            if c.get("preavvisoDisdettaGiorni"):
                y, mth, d = map(int, c["dataFine"].split("-"))
                lim = date(y, mth, d) - timedelta(days=int(c["preavvisoDisdettaGiorni"]))
                scadenze.append(
                    {
                        "id": nid("scad"),
                        "mezzoId": c["mezzoId"],
                        "tipo": "disdetta",
                        "scadenza": lim.isoformat(),
                        "descrizione": f"Ultimo giorno per disdire ({c['preavvisoDisdettaGiorni']} gg prima della fine)",
                    }
                )

    # ---------------- Assicurazioni dai PDF ----------------
    gia = {(s["mezzoId"], s["tipo"]) for s in scadenze if s.get("mezzoId")}
    for a in assicurazioni:
        targa = norm_targa(a.get("targa") or "")
        if not targa or targa not in mezzi:
            continue
        if a.get("scadenza"):
            if (targa, "assicurazione") in gia:
                # tieni la più recente: sostituisci se successiva
                for s in scadenze:
                    if s.get("mezzoId") == targa and s["tipo"] == "assicurazione":
                        if a["scadenza"] > s["scadenza"]:
                            s["scadenza"] = a["scadenza"]
                            s["descrizione"] = pulisci(a.get("compagnia"))
            else:
                scadenze.append(
                    {
                        "id": nid("scad"),
                        "mezzoId": targa,
                        "tipo": "assicurazione",
                        "scadenza": a["scadenza"],
                        "descrizione": pulisci(a.get("compagnia")),
                        "note": pulisci(a.get("numeroPolizza")),
                    }
                )
                gia.add((targa, "assicurazione"))
        if a.get("file"):
            mezzi[targa].setdefault("docAssicurazione", f"ASSICURAZIONI/{a['file']}")

    db = {
        "versione": 0,
        "mezzi": sorted(mezzi.values(), key=lambda x: x["targa"]),
        "affidamenti": affidamenti,
        "scadenze": scadenze,
        "contratti": contratti,
        "multe": [],
        "ztl": ztl,
        "officina": [],
        "km": [],
        "carburante": [],
        "parametri": {
            "emailAlert": ["d.gabelli@drlogistica.it", "f.r@drlogistica.it"],
            "soglieAlertGiorni": [90, 60, 30],
            "extraKmDefaultEur": 0.15,
            "sogliaOutlierPct": 20,
            "causaliNonConsegna": [
                "Destinatario assente",
                "Indirizzo errato",
                "Rifiuto",
                "Problemi contrassegno",
                "Causa di forza maggiore",
                "Non ho fatto in tempo",
            ],
            "costoOrarioAutistaEur": 18,
            "costoCarburanteEurLitro": 1.75,
            "consumoMedioKmLitro": 9,
        },
    }

    # Pulizia: rimuovi chiavi None
    def clean(o):
        if isinstance(o, dict):
            return {k: clean(v) for k, v in o.items() if v is not None}
        if isinstance(o, list):
            return [clean(x) for x in o]
        return o

    db = clean(db)
    out_path.write_text(json.dumps(db, ensure_ascii=False, indent=1), "utf-8")
    print(
        f"Seed scritto: {out_path}\n"
        f"  mezzi={len(db['mezzi'])} affidamenti={len(db['affidamenti'])} "
        f"scadenze={len(db['scadenze'])} contratti={len(db['contratti'])} ztl={len(db['ztl'])}"
    )
    for w in warn:
        print("  [!]", w.encode("ascii", "replace").decode())


if __name__ == "__main__":
    main()
